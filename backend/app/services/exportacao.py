"""Exportação do painel para Excel (openpyxl, com gráficos nativos) e PDF (reportlab).

Recebe a mesma `Analise` e as mesmas linhas da tela, então os números sempre conferem.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import CondPageBreak, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import NivelAlerta
from app.schemas import Analise, Grupo
from app.services.analise import ROTULO_ALERTA, ROTULO_STATUS, Linha

# Paleta validada para daltonismo (mesma ordem e cores do painel na tela). Nas barras empilhadas a ordem é
# pendente de resposta → pendente de recebimento → concluído → não confirmado (pares vizinhos validados).
COR_PENDENTE, COR_NAO_CONFIRMADO, COR_CONFIRMADO, COR_RECEBIMENTO = "2A78D6", "EB6834", "1BAF7A", "EDA100"
SITUACOES = [("Pendentes de resposta", COR_PENDENTE), ("Pendentes de recebimento", COR_RECEBIMENTO),
             ("Concluídos", COR_CONFIRMADO), ("Não confirmados", COR_NAO_CONFIRMADO)]
COR_TEXTO, COR_SUAVE, COR_GRADE = "0F172A", "64748B", "E2E8F0"
FUNDO_ALERTA = {NivelAlerta.CRITICO: "FEE2E2", NivelAlerta.ATENCAO: "FEF3C7"}
MAX_LINHAS_PDF = 2000


def _data(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _pct(v: float | None) -> str:
    return f"{v * 100:.0f}%".replace(".", ",") if v is not None else "—"


def _texto(v: int | None) -> str:
    return "" if v is None else str(v)


def _num(v: float | None) -> str:
    return f"{v:.1f}".replace(".", ",") if v is not None else "—"


def linhas_indicadores(a: Analise) -> list[tuple[str, str]]:
    i = a.indicadores
    atencao, critico = a.prazos.dias_atencao, a.prazos.dias_critico
    no_prazo, em_atencao, criticos = (f"até {atencao - 1} dias úteis", f"{atencao} a {critico - 1} dias úteis",
                                      f"{critico}+ dias úteis")
    return [
        ("Total de ofícios", str(i.total)),
        ("Pendentes de resposta", str(i.pendentes)),
        (f"   No prazo ({no_prazo})", str(i.pendentes_normal)),
        (f"   Em atenção ({em_atencao})", str(i.pendentes_atencao)),
        (f"   Críticos ({criticos})", str(i.pendentes_critico)),
        ("Pendentes de recebimento (resposta confirmada, entrega não confirmada)", str(i.pendentes_recebimento)),
        (f"   No prazo ({no_prazo} desde a resposta)", str(i.recebimento_normal)),
        (f"   Em atenção ({em_atencao})", str(i.recebimento_atencao)),
        (f"   Críticos ({criticos})", str(i.recebimento_critico)),
        ("Concluídos", str(i.concluidos)),
        ("   Recebidos (resposta e entrega confirmadas)", str(i.recebidos)),
        ("   Confirmados (sem etapa de entrega: Comunicados)", str(i.confirmados)),
        ("Não confirmados", str(i.nao_confirmados)),
        ("Taxa de confirmação", _pct(i.taxa_confirmacao)),
        ("Tempo médio de resposta (dias úteis)", _num(i.tempo_medio_resposta)),
        ("Respondidos antes de entrar em atenção", _pct(i.respondidos_no_prazo)),
        ("Tempo médio de entrega após a resposta (dias úteis)", _num(i.tempo_medio_entrega)),
        ("Entregues antes de entrar em atenção", _pct(i.entregues_no_prazo)),
    ]


CABECALHO_GRUPO = ["", "Total", "Pendentes de resposta", "Pendentes de recebimento", "Atenção", "Críticos",
                   "Confirmados", "Recebidos", "Não confirmados", "Concluídos (confirmados + recebidos)",
                   "Tempo médio de resposta (dias úteis)", "Tempo médio de entrega (dias úteis)"]
# Colunas (1 = A) das séries do gráfico empilhado, na ordem de SITUACOES
COLUNAS_SITUACAO = (3, 4, 10, 9)


def _valores_grupo(g: Grupo) -> list:
    return [g.rotulo, g.total, g.pendentes, g.pendentes_recebimento, g.atencao, g.criticos, g.confirmados,
            g.recebidos, g.nao_confirmados, g.confirmados + g.recebidos, g.tempo_medio_resposta,
            g.tempo_medio_entrega]


def _linha_oficio(linha: Linha) -> list:
    o = linha.oficio
    return [
        o.numero,
        o.data_oficio,
        linha.tipo,
        linha.destinatario,
        o.protocolo or "",
        linha.dono_nome,
        ROTULO_STATUS[o.status],
        ROTULO_ALERTA[linha.nivel] if linha.nivel else "",
        linha.dias_uteis if linha.nivel else None,
        linha.tempo_resposta,
        linha.tempo_entrega,
        o.inicio_prazo,
        o.data_resposta,
        o.data_recebimento,
        o.observacao or "",
    ]


CABECALHO_OFICIOS = [
    "Nº", "Data do ofício", "Tipo", "Destinatário", "Protocolo", "Cliente", "Status", "Prazo",
    "Dias úteis em aberto", "Dias úteis até a resposta", "Dias úteis até a entrega", "Início do prazo",
    "Data da resposta", "Data do recebimento", "Observação",
]


# ======================= Excel =======================

_borda = Border(bottom=Side(style="thin", color=COR_GRADE))
_cabecalho_fonte = Font(bold=True, color="FFFFFF")
_cabecalho_fundo = PatternFill("solid", fgColor="1D4ED8")


def _tabela(ws: Worksheet, linha: int, cabecalho: list[str], dados: list[list], larguras: list[int] | None = None) -> int:
    """Escreve uma tabela a partir de `linha`; devolve a última linha usada."""
    for c, titulo in enumerate(cabecalho, 1):
        cel = ws.cell(linha, c, titulo)
        cel.font, cel.fill = _cabecalho_fonte, _cabecalho_fundo
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    for r, valores in enumerate(dados, linha + 1):
        for c, v in enumerate(valores, 1):
            cel = ws.cell(r, c, v)
            cel.border = _borda
            if isinstance(v, date):
                cel.number_format = "dd/mm/yyyy"
    for c, largura in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = largura
    return linha + len(dados)


def _grafico_empilhado(ws_dados: Worksheet, n: int, titulo: str) -> BarChart:
    """Barras horizontais empilhadas por situação (SITUACOES, colunas COLUNAS_SITUACAO)."""
    g = BarChart()
    g.type, g.grouping, g.overlap = "bar", "stacked", 100
    g.title, g.height, g.width = titulo, 7.5, 16
    g.legend.position = "b"
    categorias = Reference(ws_dados, min_col=1, min_row=2, max_row=n + 1)
    for col, (_, cor) in zip(COLUNAS_SITUACAO, SITUACOES):
        g.add_data(Reference(ws_dados, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        g.series[-1].graphicalProperties.solidFill = cor
        g.series[-1].graphicalProperties.line.solidFill = "FFFFFF"
    g.set_categories(categorias)
    g.y_axis.majorGridlines.spPr = None
    g.x_axis.delete = False
    g.y_axis.delete = False
    return g


def gerar_excel(a: Analise, linhas: list[Linha], filtros: list[str], mostrar_clientes: bool) -> bytes:
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    resumo["A1"] = "Planner de Ofícios — Relatório"
    resumo["A1"].font = Font(bold=True, size=16, color=COR_TEXTO)
    resumo["A2"] = f"Situação em {_data(a.data_referencia)}"
    resumo["A3"] = "Filtros: " + (" · ".join(filtros) if filtros else "nenhum (todos os ofícios)")
    for c in ("A2", "A3"):
        resumo[c].font = Font(color=COR_SUAVE)
    _tabela(resumo, 5, ["Indicador", "Valor"], [list(x) for x in linhas_indicadores(a)], [62, 14])

    abas: list[tuple[str, list[Grupo]]] = [
        ("Por tipo", a.por_tipo),
        ("Por destinatário", a.por_destinatario),
        ("Principais setores", a.por_setor),
    ]
    if mostrar_clientes:
        abas.append(("Por cliente", a.por_cliente))
    planilhas = {}
    for nome, grupos in abas:
        ws = wb.create_sheet(nome)
        _tabela(ws, 1, CABECALHO_GRUPO, [_valores_grupo(g) for g in grupos],
                [34, 9, 13, 14, 10, 10, 13, 12, 14, 18, 16, 16])
        ws.cell(1, 1, nome.replace("Por ", "").capitalize())
        ws.freeze_panes = "B2"
        planilhas[nome] = (ws, len(grupos))

    meses = wb.create_sheet("Por mês")
    _tabela(meses, 1, ["Mês", "Cadastrados", "Respondidos", "Recebidos"],
            [[f"{m.mes[5:]}/{m.mes[:4]}", m.cadastrados, m.respondidos, m.recebidos] for m in a.por_mes],
            [12, 13, 13, 13])

    oficios = wb.create_sheet("Ofícios")
    fim = _tabela(oficios, 1, CABECALHO_OFICIOS, [_linha_oficio(x) for x in linhas],
                  [12, 13, 26, 30, 18, 20, 24, 11, 12, 13, 13, 14, 15, 15, 60])
    oficios.freeze_panes = "B2"
    oficios.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO_OFICIOS))}{max(fim, 1)}"
    for r, x in enumerate(linhas, 2):
        if x.nivel in FUNDO_ALERTA:
            oficios.cell(r, 8).fill = PatternFill("solid", fgColor=FUNDO_ALERTA[x.nivel])
        oficios.cell(r, len(CABECALHO_OFICIOS)).alignment = Alignment(wrap_text=True, vertical="top")

    # Gráficos nativos do Excel na aba Resumo (continuam editáveis no Excel)
    ws_tipo, n_tipo = planilhas["Por tipo"]
    if n_tipo:
        resumo.add_chart(_grafico_empilhado(ws_tipo, n_tipo, "Ofícios por tipo"), "D5")
    ws_dest, n_dest = planilhas["Por destinatário"]
    if n_dest:
        resumo.add_chart(_grafico_empilhado(ws_dest, n_dest, "Ofícios por destinatário"), "D21")
    if a.por_mes:
        linha_mes = LineChart()
        linha_mes.title, linha_mes.height, linha_mes.width = "Evolução mensal", 7.5, 16
        linha_mes.legend.position = "b"
        linha_mes.add_data(Reference(meses, min_col=2, max_col=4, min_row=1, max_row=len(a.por_mes) + 1),
                           titles_from_data=True)
        linha_mes.set_categories(Reference(meses, min_col=1, min_row=2, max_row=len(a.por_mes) + 1))
        for serie, cor in zip(linha_mes.series, (COR_PENDENTE, COR_NAO_CONFIRMADO, COR_CONFIRMADO)):
            serie.graphicalProperties.line.solidFill = cor
            serie.graphicalProperties.line.width = 22000
            serie.smooth = False
        linha_mes.x_axis.delete = False
        linha_mes.y_axis.delete = False
        resumo.add_chart(linha_mes, "D37")

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ======================= PDF =======================

def _cor(hexa: str) -> colors.Color:
    return colors.HexColor(f"#{hexa}")


def _grafico_pdf(grupos: list[Grupo], titulo: str, largura: float, altura: float) -> Drawing:
    d = Drawing(largura, altura)
    d.add(String(0, altura - 12, titulo, fontName="Helvetica-Bold", fontSize=10, fillColor=_cor(COR_TEXTO)))
    if not grupos:
        d.add(String(0, altura - 30, "Sem dados", fontName="Helvetica", fontSize=9, fillColor=_cor(COR_SUAVE)))
        return d
    g = HorizontalBarChart()
    g.x, g.y = 175, 44  # espaço para rótulos longos (ex.: Peças para Manutenção/Engenharia Clínica)
    g.width, g.height = largura - 195, altura - 76
    g.data = [
        [x.pendentes for x in grupos],
        [x.pendentes_recebimento for x in grupos],
        [x.confirmados + x.recebidos for x in grupos],
        [x.nao_confirmados for x in grupos],
    ]
    g.categoryAxis.categoryNames = [x.rotulo for x in grupos]
    g.categoryAxis.style = "stacked"
    g.categoryAxis.labels.fontName, g.categoryAxis.labels.fontSize = "Helvetica", 7.5
    g.categoryAxis.labels.boxAnchor = "e"
    g.categoryAxis.strokeColor = _cor(COR_GRADE)
    g.categoryAxis.reverseDirection = True
    g.valueAxis.valueMin = 0
    g.valueAxis.labels.fontSize = 7
    g.valueAxis.strokeColor = _cor(COR_GRADE)
    g.valueAxis.visibleGrid, g.valueAxis.gridStrokeColor = True, _cor(COR_GRADE)
    maior = max((x.total for x in grupos), default=0)
    g.valueAxis.valueMax = max(4, maior)
    g.valueAxis.valueStep = max(1, round(g.valueAxis.valueMax / 4))
    g.barWidth, g.groupSpacing = 8, 6
    for i, (_, cor) in enumerate(SITUACOES):
        g.bars[i].fillColor = _cor(cor)
        g.bars[i].strokeColor = colors.white
        g.bars[i].strokeWidth = 1
    d.add(g)
    d.add(_legenda(175, 26, SITUACOES, linhas=2, deltax=110))
    return d


def _legenda(x: float, y: float, itens: list[tuple[str, str]], linhas: int = 1, deltax: float = 85) -> Legend:
    leg = Legend()
    leg.x, leg.y = x, y
    leg.alignment, leg.columnMaximum = "right", linhas
    leg.fontName, leg.fontSize = "Helvetica", 7.5
    leg.dx = leg.dy = 7
    leg.deltax = deltax
    leg.colorNamePairs = [(_cor(c), n) for n, c in itens]
    return leg


def _grafico_mensal(a: Analise, largura: float, altura: float) -> Drawing:
    d = Drawing(largura, altura)
    d.add(String(0, altura - 12, "Evolução mensal", fontName="Helvetica-Bold", fontSize=10, fillColor=_cor(COR_TEXTO)))
    if not a.por_mes:
        return d
    g = VerticalBarChart()
    g.x, g.y = 30, 56
    g.width, g.height = largura - 40, altura - 88
    g.data = [[m.cadastrados for m in a.por_mes], [m.respondidos for m in a.por_mes],
              [m.recebidos for m in a.por_mes]]
    g.categoryAxis.categoryNames = [f"{m.mes[5:]}/{m.mes[2:4]}" for m in a.por_mes]
    g.categoryAxis.labels.fontSize = 7
    g.categoryAxis.labels.angle = 45 if len(a.por_mes) > 8 else 0
    g.categoryAxis.labels.boxAnchor = "ne" if len(a.por_mes) > 8 else "n"
    g.categoryAxis.strokeColor = _cor(COR_GRADE)
    g.valueAxis.valueMin = 0
    maior = max([max(m.cadastrados, m.respondidos, m.recebidos) for m in a.por_mes] + [4])
    g.valueAxis.valueMax = maior
    g.valueAxis.valueStep = max(1, round(maior / 4))
    g.valueAxis.labels.fontSize = 7
    g.valueAxis.strokeColor = _cor(COR_GRADE)
    g.valueAxis.visibleGrid, g.valueAxis.gridStrokeColor = True, _cor(COR_GRADE)
    g.barSpacing, g.groupSpacing = 1, 6
    mensal = [("Cadastrados (pela data do ofício)", COR_PENDENTE),
              ("Respondidos (pela data da resposta)", COR_NAO_CONFIRMADO),
              ("Recebidos (pela data do recebimento)", COR_CONFIRMADO)]
    for i, (_, cor) in enumerate(mensal):
        g.bars[i].fillColor, g.bars[i].strokeColor = _cor(cor), colors.white
    d.add(g)
    d.add(_legenda(30, 20, mensal, linhas=2, deltax=150))
    return d


def _estilo_tabela(numerica_desde: int = 1) -> TableStyle:
    return TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), _cor("1D4ED8")),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, _cor(COR_GRADE)),
        ("ALIGN", (numerica_desde, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])


def _tabela_grupos(titulo: str, grupos: list[Grupo]) -> list:
    if not grupos:
        return []
    cab = ParagraphStyle("cab", fontName="Helvetica-Bold", fontSize=7, leading=8.5, textColor=colors.white,
                         alignment=2)
    titulos = ["Total", "Pend. resposta", "Pend. recebimento", "Atenção", "Críticos", "Confirmados", "Recebidos",
               "Não confirmados", "Tempo médio de resposta", "Tempo médio de entrega"]
    dados = [[titulo] + [Paragraph(t, cab) for t in titulos]] + [
        [g.rotulo, g.total, g.pendentes, g.pendentes_recebimento, g.atencao, g.criticos, g.confirmados, g.recebidos,
         g.nao_confirmados, _num(g.tempo_medio_resposta), _num(g.tempo_medio_entrega)]
        for g in grupos
    ]
    t = Table(dados, colWidths=[59 * mm] + [21 * mm] * len(titulos), repeatRows=1)
    t.setStyle(_estilo_tabela())
    return [KeepTogether([t]), Spacer(1, 6 * mm)]


def gerar_pdf(a: Analise, linhas: list[Linha], filtros: list[str], mostrar_clientes: bool) -> bytes:
    saida = BytesIO()
    pagina = landscape(A4)
    doc = SimpleDocTemplate(saida, pagesize=pagina, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title="Planner de Ofícios — Relatório")
    base = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=base["Title"], fontName="Helvetica-Bold", fontSize=18, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=4)
    suave = ParagraphStyle("suave", parent=base["Normal"], fontName="Helvetica", fontSize=9,
                           textColor=_cor(COR_SUAVE))
    celula = ParagraphStyle("celula", parent=base["Normal"], fontName="Helvetica", fontSize=7, leading=8.5)
    celula_cab = ParagraphStyle("celula_cab", parent=celula, fontName="Helvetica-Bold", fontSize=7.5,
                                textColor=colors.white)
    util = pagina[0] - 28 * mm

    conteudo: list = [
        Paragraph("Planner de Ofícios — Relatório", h1),
        Paragraph(f"Situação em {_data(a.data_referencia)} · Filtros: "
                  + (" · ".join(filtros) if filtros else "nenhum (todos os ofícios)"), suave),
        Spacer(1, 5 * mm),
    ]

    # Indicadores em blocos
    i = a.indicadores
    rotulo_bloco = ParagraphStyle("bloco", fontName="Helvetica", fontSize=7.5, leading=9, textColor=_cor(COR_SUAVE))
    destaque = [
        (str(i.total), "Total de ofícios"),
        (str(i.pendentes), "Pendentes de resposta"),
        (str(i.pendentes_recebimento), "Pendentes de recebimento"),
        (str(i.pendentes_critico + i.recebimento_critico), f"Críticos ({a.prazos.dias_critico}+ dias úteis)"),
        (str(i.concluidos), "Concluídos"),
        (str(i.nao_confirmados), "Não confirmados"),
        (_pct(i.taxa_confirmacao), "Taxa de confirmação"),
        (_num(i.tempo_medio_resposta), "Tempo médio de resposta (dias úteis)"),
        (_num(i.tempo_medio_entrega), "Tempo médio de entrega (dias úteis)"),
    ]
    blocos = Table(
        [[v for v, _ in destaque], [Paragraph(r, rotulo_bloco) for _, r in destaque]],
        colWidths=[util / len(destaque)] * len(destaque),
    )
    blocos.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 16),
        ("TEXTCOLOR", (3, 0), (3, 0), _cor("B91C1C")),
        ("VALIGN", (0, 1), (-1, 1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.5, _cor(COR_GRADE)),
        ("LINEAFTER", (0, 0), (-2, -1), 0.5, _cor(COR_GRADE)),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
    ]))
    conteudo += [blocos, Spacer(1, 6 * mm)]

    metade = util / 2 - 4 * mm
    altura_tipo = 84 + 16 * max(len(a.por_tipo), 1)
    graficos = Table([[_grafico_pdf(a.por_tipo, "Ofícios por tipo", metade, altura_tipo),
                       _grafico_mensal(a, metade, altura_tipo)]], colWidths=[metade + 8 * mm, metade])
    graficos.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    conteudo += [graficos, Spacer(1, 4 * mm)]
    conteudo += [_grafico_pdf(a.por_destinatario, "Ofícios por destinatário", metade,
                              84 + 16 * max(len(a.por_destinatario), 1))]

    # Quebras condicionais: só mudam de página quando falta espaço (evita páginas quase vazias)
    conteudo += [Spacer(1, 4 * mm), CondPageBreak(70 * mm), Paragraph("Detalhamento", h2)]
    conteudo += _tabela_grupos("Tipo", a.por_tipo)
    conteudo += _tabela_grupos("Destinatário", a.por_destinatario)
    conteudo += _tabela_grupos("Principais setores / destinos", a.por_setor)
    if mostrar_clientes:
        conteudo += _tabela_grupos("Cliente", a.por_cliente)

    # Lista de ofícios
    conteudo += [CondPageBreak(90 * mm), Paragraph(f"Ofícios ({len(linhas)})", h2)]
    if len(linhas) > MAX_LINHAS_PDF:
        conteudo.append(Paragraph(f"Mostrando os {MAX_LINHAS_PDF} mais recentes; a lista completa está no Excel.", suave))
    cab = ["Nº", "Data", "Tipo", "Destinatário", "Protocolo", "Cliente", "Status", "Prazo",
           Paragraph("Dias em aberto", celula_cab), Paragraph("Dias até resposta", celula_cab),
           Paragraph("Dias até entrega", celula_cab), "Resposta", "Receb.", "Observação"]
    dados = [cab]
    estilo = _estilo_tabela(numerica_desde=99)
    for i, x in enumerate(linhas[:MAX_LINHAS_PDF], 1):
        o = x.oficio
        dados.append([
            Paragraph(o.numero, celula), _data(o.data_oficio),
            Paragraph(x.tipo, celula),
            Paragraph(x.destinatario, celula), Paragraph(o.protocolo or "", celula), Paragraph(x.dono_nome, celula),
            Paragraph(ROTULO_STATUS[o.status], celula), ROTULO_ALERTA[x.nivel] if x.nivel else "",
            str(x.dias_uteis) if x.nivel else "", _texto(x.tempo_resposta), _texto(x.tempo_entrega),
            _data(o.data_resposta), _data(o.data_recebimento),
            Paragraph(o.observacao or "", celula),
        ])
        if x.nivel in FUNDO_ALERTA:
            estilo.add("BACKGROUND", (7, i), (7, i), _cor(FUNDO_ALERTA[x.nivel]))
        if x.nivel == NivelAlerta.CRITICO:
            estilo.add("TEXTCOLOR", (7, i), (7, i), _cor("B91C1C"))
    larguras = [18, 17, 24, 28, 18, 20, 24, 15, 12, 12, 12, 17, 17, 0]
    larguras[-1] = util / mm - sum(larguras)
    tabela = Table(dados, colWidths=[w * mm for w in larguras], repeatRows=1)
    tabela.setStyle(estilo)
    conteudo.append(tabela)

    def rodape(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(_cor(COR_SUAVE))
        canvas.drawString(14 * mm, 8 * mm, f"Planner de Ofícios · gerado em {_data(a.data_referencia)}")
        canvas.drawRightString(pagina[0] - 14 * mm, 8 * mm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(conteudo, onFirstPage=rodape, onLaterPages=rodape)
    return saida.getvalue()

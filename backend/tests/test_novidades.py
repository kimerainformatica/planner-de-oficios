"""Destinatário, tipo, protocolo opcional, evidência da resposta, análise do painel e exportações."""

from io import BytesIO

from openpyxl import load_workbook

PDF = b"%PDF-1.4 evidencia"


def criar(api, quem, numero, data="2026-10-01", **extra):
    corpo = {"numero": numero, "data_oficio": data, "tipo": "EPIS", **extra}
    return api.post("/api/oficios", json=corpo, headers=quem)


def responder(api, quem, oid, status, data_resposta=None):
    return api.patch(f"/api/oficios/{oid}/status", json={"status": status, "data_resposta": data_resposta},
                     headers=quem)


# ---------- cadastro ----------


def test_protocolo_opcional(api, ana):
    assert criar(api, ana, "1").json()["protocolo"] is None
    assert criar(api, ana, "2", protocolo="  ").json()["protocolo"] is None  # vazio vira nulo
    assert criar(api, ana, "3").status_code == 201  # vários sem protocolo convivem


def test_tipo_obrigatorio_no_cadastro(api, ana):
    r = api.post("/api/oficios", json={"numero": "1", "data_oficio": "2026-10-01"}, headers=ana)
    assert r.status_code == 422
    assert criar(api, ana, "2", tipo="INVALIDO").status_code == 422


def test_destinatario(api, ana):
    r = criar(api, ana, "1", destinatario_tipo="SES", destinatario_setor="Almoxarifado central")
    assert (r.json()["destinatario_tipo"], r.json()["destinatario_setor"]) == ("SES", "Almoxarifado central")
    assert criar(api, ana, "2", destinatario_tipo="ACQUA").status_code == 422  # falta o setor
    assert criar(api, ana, "3", destinatario_setor="Sem tipo").status_code == 422  # falta o tipo

    oid = r.json()["id"]
    r = api.patch(f"/api/oficios/{oid}", json={"destinatario_tipo": "OUTROS", "destinatario_setor": "Prefeitura"},
                  headers=ana)
    assert (r.json()["destinatario_tipo"], r.json()["destinatario_setor"]) == ("OUTROS", "Prefeitura")
    # Trocar só o tipo mantém o setor; apagar exige apagar os dois
    assert api.patch(f"/api/oficios/{oid}", json={"destinatario_tipo": None}, headers=ana).status_code == 422
    r = api.patch(f"/api/oficios/{oid}", json={"destinatario_tipo": None, "destinatario_setor": None}, headers=ana)
    assert r.json()["destinatario_tipo"] is None
    r = api.patch(f"/api/oficios/{oid}", json={"tipo": "SERVICOS", "protocolo": None}, headers=ana)
    assert (r.json()["tipo"], r.json()["protocolo"]) == ("SERVICOS", None)


def test_filtros_por_tipo_destinatario_e_busca(api, ana):
    criar(api, ana, "1", tipo="EPIS", destinatario_tipo="SES", destinatario_setor="Farmácia")
    criar(api, ana, "2", tipo="SERVICOS", destinatario_tipo="ACQUA", destinatario_setor="Manutenção",
          observacao="Troca de ar-condicionado")
    criar(api, ana, "3", tipo="SERVICOS")

    def numeros(**params):
        return sorted(o["numero"] for o in api.get("/api/oficios", params=params, headers=ana).json()["itens"])

    assert numeros(tipo="SERVICOS") == ["2", "3"]
    assert numeros(destinatario_tipo="SES") == ["1"]
    assert numeros(busca="farmácia") == ["1"]  # busca no setor
    assert numeros(busca="ar-condicionado") == ["2"]  # busca na observação


# ---------- evidência da resposta ----------


def test_evidencia_da_resposta(api, admin, ana, bruno):
    oid = criar(api, ana, "1").json()["id"]
    url = f"/api/oficios/{oid}/anexos/RESPOSTA"
    arquivo = {"arquivo": ("resposta.pdf", PDF, "application/pdf")}

    # Sem resposta registrada não há onde ligar a evidência
    assert api.post(url, files=arquivo, headers=ana).status_code == 400

    r = responder(api, ana, oid, "NAO_CONFIRMADO", "2026-10-05").json()
    evento1 = r["historico"][-1]["id"]
    assert api.post(url, params={"historico_id": evento1}, files=arquivo, headers=ana).status_code == 201
    assert api.post(url, params={"historico_id": r["historico"][0]["id"]}, files=arquivo,
                    headers=ana).status_code == 400  # evento de cadastro não é resposta
    assert api.post(url, params={"historico_id": evento1}, files=arquivo, headers=bruno).status_code == 404

    responder(api, ana, oid, "PENDENTE_RESPOSTA")
    r = responder(api, ana, oid, "CONFIRMADO", "2026-10-20").json()
    evento2 = r["historico"][-1]["id"]
    # Confirmado: o dono ainda pode anexar a evidência (é exatamente o momento de anexá-la)
    nova = {"arquivo": ("confirmacao.pdf", PDF + b"2", "application/pdf")}
    assert api.post(url, params={"historico_id": evento2}, files=nova, headers=ana).status_code == 201

    detalhe = api.get(f"/api/oficios/{oid}", headers=ana).json()
    evidencias = {h["id"]: (h["evidencia"] or {}).get("nome_original") for h in detalhe["historico"]}
    assert evidencias[evento1] == "resposta.pdf"  # a evidência do ciclo anterior continua no histórico
    assert evidencias[evento2] == "confirmacao.pdf"
    assert sum(a["tipo"] == "RESPOSTA" for a in detalhe["anexos"]) == 2

    doc = {"arquivo": ("o.pdf", PDF, "application/pdf")}
    # Reenviar a evidência do mesmo evento substitui
    assert api.post(url, params={"historico_id": evento2}, files=doc, headers=admin).status_code == 201
    detalhe = api.get(f"/api/oficios/{oid}", headers=ana).json()
    assert sum(a["tipo"] == "RESPOSTA" for a in detalhe["anexos"]) == 2


# ---------- painel e exportações ----------


def _cenario(api, admin, ana, bruno):
    a1 = criar(api, ana, "1", data="2026-08-03", tipo="EPIS", destinatario_tipo="SES", destinatario_setor="Farmácia")
    criar(api, ana, "2", data="2026-10-14", tipo="EPIS", destinatario_tipo="SES", destinatario_setor="farmácia ")
    b1 = criar(api, bruno, "1", data="2026-09-01", tipo="SERVICOS", destinatario_tipo="ACQUA",
               destinatario_setor="Obras")
    responder(api, admin, a1.json()["id"], "CONFIRMADO", "2026-08-10")  # 5 dias úteis
    responder(api, admin, b1.json()["id"], "NAO_CONFIRMADO", "2026-09-22")  # 15 dias úteis


def test_analise_do_painel(api, admin, ana, bruno):
    _cenario(api, admin, ana, bruno)
    a = api.get("/api/dashboard/analise", headers=admin).json()
    i = a["indicadores"]
    # EPI confirmado aguarda o recebimento: conta como resposta confirmada, não como concluído
    assert (i["total"], i["pendentes"], i["pendentes_recebimento"], i["confirmados"], i["nao_confirmados"]) == (
        3, 1, 1, 0, 1)
    assert i["taxa_confirmacao"] == 0.5
    assert i["tempo_medio_resposta"] == 10.0  # (5 + 15) / 2
    assert i["respondidos_no_prazo"] == 0.5  # só o de 5 dias ficou abaixo de 15

    tipos = {g["chave"]: g for g in a["por_tipo"]}
    assert len(tipos) == 8  # todos os tipos aparecem, mesmo zerados
    assert (tipos["EPIS"]["total"], tipos["EPIS"]["pendentes_recebimento"], tipos["EPIS"]["pendentes"]) == (2, 1, 1)
    assert tipos["SERVICOS"]["nao_confirmados"] == 1
    setores = {g["rotulo"]: g["total"] for g in a["por_setor"]}
    assert setores["SES · Farmácia"] == 2  # agrupa ignorando maiúsculas/espaços
    assert {g["rotulo"] for g in a["por_cliente"]} == {"Ana", "Bruno"}
    meses = {m["mes"]: m for m in a["por_mes"]}
    assert meses["2026-08"] == {"mes": "2026-08", "cadastrados": 1, "respondidos": 1, "recebidos": 0}
    assert meses["2026-10"]["cadastrados"] == 1

    # Filtros valem para tudo; cliente vê só os seus e não recebe o quadro por cliente
    f = api.get("/api/dashboard/analise", params={"tipo": "SERVICOS"}, headers=admin).json()
    assert f["indicadores"]["total"] == 1
    da_ana = api.get("/api/dashboard/analise", headers=ana).json()
    assert da_ana["indicadores"]["total"] == 2 and da_ana["por_cliente"] == []


def test_exportar_excel_e_pdf(api, admin, ana, bruno):
    _cenario(api, admin, ana, bruno)
    r = api.get("/api/dashboard/exportar", params={"formato": "xlsx", "tipo": "EPIS"}, headers=admin)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/octet-stream"
    assert "oficios-" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Resumo", "Por tipo", "Por destinatário", "Principais setores", "Por cliente",
                             "Por mês", "Ofícios"]
    assert "Tipo: EPI's" in wb["Resumo"]["A3"].value
    assert wb["Ofícios"].max_row == 3  # cabeçalho + 2 ofícios de EPI
    assert len(wb["Resumo"]._charts) == 3

    r = api.get("/api/dashboard/exportar", params={"formato": "pdf"}, headers=ana)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert api.get("/api/dashboard/exportar", params={"formato": "doc"}, headers=ana).status_code == 422
